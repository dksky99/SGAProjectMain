from pathlib import Path
import json
import re
from PIL import Image, ImageDraw, ImageFont
from docx import Document
from docx.shared import Cm, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_TAB_ALIGNMENT, WD_TAB_LEADER
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

ROOT = Path(__file__).resolve().parents[2]
NOTES = Path(__file__).resolve().parent
WORK = Path('C:/Users/Administrator/.codex/visualizations/2026/10/09/01a12173-52ea-7e52-979d-1bc5d532652f/implementation_manual_work')
WORK.mkdir(parents=True, exist_ok=True)
OUT = ROOT / 'docs/SGAProjectMain_Implementation_Manual_2026-10-10.docx'
SOURCES = ['00_overview.md','01_character_animation_damage.md','02_ai_world_mission_objects.md','04_equipment_stratagem.md','03_game_flow_ui_save_data.md','05_reading_and_change_map.md']
CHAPTERS = ['프로젝트 전체 구조와 읽는 순서','캐릭터 애니메이션 상태와 부위 피해','AI 적 분대와 월드 임무 오브젝트','장비 투사체와 스트라타젬','게임 진행 UI 세이브와 데이터','실행 시나리오와 변경 영향 찾아보기']
FONT = '맑은 고딕'
BLACK = RGBColor(0,0,0)

def clean_heading(text):
    text = re.sub(r'^\d+(?:\.\d+)*[.)]?\s*', '', text)
    text = text.replace('·',' ').replace(',',' ').replace(':',' ').replace('/',' ').replace('→',' ').replace('—',' ').replace('–',' ')
    text = re.sub(r'[()\[\]`*]', '', text)
    return re.sub(r'\s+', ' ', text).strip()

def set_font(run, size=None, bold=None, name=FONT):
    run.font.name = name
    run.font.color.rgb = BLACK
    if size is not None: run.font.size = Pt(size)
    if bold is not None: run.bold = bold
    rpr = run._element.get_or_add_rPr()
    fonts = rpr.find(qn('w:rFonts'))
    if fonts is None:
        fonts = OxmlElement('w:rFonts'); rpr.append(fonts)
    for attr in ('ascii','hAnsi','eastAsia','cs'): fonts.set(qn('w:'+attr), name)
    lang = OxmlElement('w:lang'); lang.set(qn('w:val'),'ko-KR'); lang.set(qn('w:eastAsia'),'ko-KR'); rpr.append(lang)

def inline(p, text, size=None):
    parts = re.split(r'(`[^`]+`|\*\*[^*]+\*\*)', text)
    for part in parts:
        if not part: continue
        is_code = part.startswith('`') and part.endswith('`')
        is_bold = part.startswith('**') and part.endswith('**')
        value = part[1:-1] if is_code else part[2:-2] if is_bold else part
        if is_code:
            value = value.replace('/', '/\u200b')
        run=p.add_run(value)
        set_font(run, size=(size or 9.4) if is_code else size, bold=True if is_bold else None, name='Consolas' if is_code else FONT)
        if is_code:
            run._element.rPr.rFonts.set(qn('w:eastAsia'),FONT)

def bookmark(p, anchor, index):
    start=OxmlElement('w:bookmarkStart'); start.set(qn('w:id'),str(index)); start.set(qn('w:name'),anchor)
    end=OxmlElement('w:bookmarkEnd'); end.set(qn('w:id'),str(index))
    p._p.insert(0,start); p._p.append(end)

def toc_link(p, text, anchor, size, bold):
    link=OxmlElement('w:hyperlink'); link.set(qn('w:anchor'),anchor)
    r=OxmlElement('w:r'); prop=OxmlElement('w:rPr')
    fonts=OxmlElement('w:rFonts')
    for attr in ('ascii','hAnsi','eastAsia'): fonts.set(qn('w:'+attr),FONT)
    prop.append(fonts)
    col=OxmlElement('w:color'); col.set(qn('w:val'),'000000'); prop.append(col)
    sz=OxmlElement('w:sz'); sz.set(qn('w:val'),str(int(size*2))); prop.append(sz)
    if bold: prop.append(OxmlElement('w:b'))
    r.append(prop); t=OxmlElement('w:t'); t.text=text; r.append(t); link.append(r); p._p.append(link)

def table_border(table):
    pr=table._tbl.tblPr
    borders=OxmlElement('w:tblBorders')
    for edge in ('top','left','bottom','right','insideH','insideV'):
        e=OxmlElement('w:'+edge); e.set(qn('w:val'),'single'); e.set(qn('w:sz'),'5'); e.set(qn('w:color'),'D9D9D9'); borders.append(e)
    pr.append(borders)

def make_table(doc, rows):
    n=len(rows[0])
    table=doc.add_table(rows=0, cols=n)
    table.alignment=WD_TABLE_ALIGNMENT.CENTER
    table.autofit=False
    ratios={2:[0.31,0.69],3:[0.25,0.28,0.47],4:[0.17,0.23,0.29,0.31]}.get(n,[1/n]*n)
    widths=[17*r for r in ratios]
    for c,w in zip(table.columns,widths): c.width=Cm(w)
    table_border(table)
    vertical_padding=55 if rows[0][0]=='용어' else 85
    for ri,rowdata in enumerate(rows):
        row=table.add_row()
        trpr=row._tr.get_or_add_trPr()
        trpr.append(OxmlElement('w:cantSplit'))
        if ri==0:
            repeat=OxmlElement('w:tblHeader'); repeat.set(qn('w:val'),'true'); trpr.append(repeat)
        for ci,cell in enumerate(row.cells):
            cell.width=Cm(widths[ci]); cell.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
            tcpr=cell._tc.get_or_add_tcPr()
            margins=OxmlElement('w:tcMar')
            for edge,val in (('top',vertical_padding),('bottom',vertical_padding),('left',100),('right',100)):
                e=OxmlElement('w:'+edge); e.set(qn('w:w'),str(val)); e.set(qn('w:type'),'dxa'); margins.append(e)
            tcpr.append(margins)
            shading=OxmlElement('w:shd'); shading.set(qn('w:fill'),'243746' if ri==0 else 'F2F5F7' if ri%2==0 else 'FFFFFF'); tcpr.append(shading)
            p=cell.paragraphs[0]
            p.paragraph_format.space_after=Pt(0); p.paragraph_format.space_before=Pt(0); p.paragraph_format.line_spacing=1.12
            p.paragraph_format.keep_with_next=False
            inline(p,rowdata[ci] if ci<len(rowdata) else '',size=9.0)
            if ri==0:
                for run in p.runs: run.bold=True; run.font.color.rgb=RGBColor(255,255,255)
    spacer=doc.add_paragraph(); spacer.paragraph_format.space_after=Pt(4); spacer.paragraph_format.space_before=Pt(0)
    spacer.paragraph_format.line_spacing=0.35
    return table

def diagram(name, nodes, arrows, height=600):
    path=WORK/(name+'.png')
    im=Image.new('RGB',(1500,height),'white'); draw=ImageDraw.Draw(im)
    font=ImageFont.truetype('C:/Windows/Fonts/malgun.ttf',27)
    small=ImageFont.truetype('C:/Windows/Fonts/malgun.ttf',23)
    for a,b in arrows:
        x1,y1=a; x2,y2=b
        draw.line((x1,y1,x2,y2),fill='#526978',width=4)
        import math
        angle=math.atan2(y2-y1,x2-x1)
        p1=(x2-17*math.cos(angle-.45),y2-17*math.sin(angle-.45))
        p2=(x2-17*math.cos(angle+.45),y2-17*math.sin(angle+.45))
        draw.polygon([(x2,y2),p1,p2],fill='#526978')
    for box,label in nodes:
        draw.rounded_rectangle(box, radius=12, fill='#EFF4F7',outline='#8397A3',width=2)
        lines=label.split('\n'); x1,y1,x2,y2=box
        y=(y1+y2)/2-len(lines)*19
        for line in lines:
            f=small if len(line)>24 else font
            bounds=draw.textbbox((0,0),line,font=f)
            draw.text(((x1+x2)/2-(bounds[2]-bounds[0])/2,y),line,font=f,fill='black')
            y+=39
    im.save(path)
    return path

fig_overview=diagram('flow_overview',[
 ((30,40,430,170),'공용 정의와 영속 기록\nGameInstance  SaveGame'),
 ((550,40,950,170),'로비와 출격 선택\nPreDeploymentState'),
 ((1070,40,1470,170),'전투 맵 진입\nOpenLevel  GameMode'),
 ((1070,330,1470,460),'캐릭터 장비 적 AI\n행동  피격  목표 수행'),
 ((550,330,950,460),'탈출과 임무 정산\nMissionResult'),
 ((30,330,430,460),'기록 저장과 로비 복귀\n다음 선택과 재실행')],
 [((430,105),(550,105)),((950,105),(1070,105)),((1270,170),(1270,330)),((1070,395),(950,395)),((550,395),(430,395)),((230,330),(230,170))],500)
fig_character=diagram('flow_character',[
 ((40,40,420,170),'입력과 화면 문맥\nPlayerCharacter'),
 ((560,40,940,170),'행동과 실제 상태\nHellDiver  State  Stat'),
 ((1080,40,1460,170),'애니메이션 입력 값\nCharacterAnimInstance'),
 ((1080,330,1460,460),'시퀀스와 재생 계산\nAnimLayer'),
 ((560,330,940,460),'시점과 완료 알림\nMontage  AnimNotify'),
 ((40,330,420,460),'행동 적용과 복구\n이동  공격  모션 종료')],
 [((420,105),(560,105)),((940,105),(1080,105)),((1270,170),(1270,330)),((1080,395),(940,395)),((560,395),(420,395)),((230,330),(230,170))],500)
fig_ai=diagram('flow_ai',[
 ((20,40,420,170),'시각 청각 피해 자극\nEnemyController'),
 ((550,40,950,170),'경계 타겟 전투 판단\nBlackboard  Service'),
 ((1080,40,1480,170),'행동 시작과 이동\nBT Task  Enemy'),
 ((1080,330,1480,460),'공격 유효 시점과 종료\nMontage  Notify'),
 ((550,330,950,460),'개별 유닛 실행 수명\nSpawn  Reset  Possess'),
 ((20,330,420,460),'분대 풀과 호출 정책\nSquad  Spawner  Manager')],
 [((420,105),(550,105)),((950,105),(1080,105)),((1280,170),(1280,330)),((1280,330),(1280,170)),((420,395),(550,395)),((950,395),(1080,395))],500)
fig_gear=diagram('flow_equipment',[
 ((20,40,420,170),'출격 로드아웃 ID\nState  DataTable'),
 ((550,40,950,170),'전장 장비와 사용 판단\nInven  Gun  Stratagem'),
 ((1080,40,1480,170),'발사 또는 투척\nFire  Throwable'),
 ((1080,330,1480,460),'탄 포드 지원 Controller\nProjectile  Deploy'),
 ((550,330,950,460),'피격과 부위 파괴\nDamageEvent  Stat'),
 ((20,330,420,460),'자원과 화면 갱신\nAmmo  Cooldown  HUD')],
 [((420,105),(550,105)),((950,105),(1080,105)),((1280,170),(1280,330)),((1080,395),(950,395)),((550,395),(420,395))],500)

doc=Document()
sec=doc.sections[0]
sec.page_width=Cm(21); sec.page_height=Cm(29.7)
sec.top_margin=Cm(1.9); sec.bottom_margin=Cm(1.8); sec.left_margin=Cm(2); sec.right_margin=Cm(2)
sec.footer_distance=Cm(.85)
for sty in ['Normal','Title','Subtitle','Heading 1','Heading 2','Heading 3','List Bullet','List Number']:
    s=doc.styles[sty]; s.font.name=FONT; s.font.color.rgb=BLACK; s.font.italic=False
    if s.element.pPr is not None:
        for border in list(s.element.pPr.findall(qn('w:pBdr'))): s.element.pPr.remove(border)
    if s.element.rPr is not None:
        for color in s.element.rPr.findall(qn('w:color')):
            for attr in ('themeColor','themeTint','themeShade'):
                color.attrib.pop(qn('w:'+attr),None)
    s.element.get_or_add_rPr().append(OxmlElement('w:rFonts'))
    fonts=s.element.rPr.findall(qn('w:rFonts'))[-1]
    for attr in ('ascii','hAnsi','eastAsia','cs'): fonts.set(qn('w:'+attr),FONT)
doc.styles['Normal'].font.size=Pt(10.5)
nf=doc.styles['Normal'].paragraph_format
nf.line_spacing=1.22; nf.space_after=Pt(6); nf.widow_control=True
for sty,size,before,after in [('Title',22,0,16),('Subtitle',12,0,9),('Heading 1',17,0,12),('Heading 2',12.8,14,7),('Heading 3',11.2,10,5)]:
    s=doc.styles[sty]; s.font.size=Pt(size); s.font.bold=sty not in ['Subtitle']
    s.paragraph_format.space_before=Pt(before); s.paragraph_format.space_after=Pt(after)
    s.paragraph_format.keep_with_next=True
    s.paragraph_format.keep_together=True
doc.styles['Heading 1'].paragraph_format.page_break_before=True
doc.styles['List Bullet'].paragraph_format.space_after=Pt(4)
doc.styles['List Number'].paragraph_format.space_after=Pt(4)

footer=sec.footer.paragraphs[0]; footer.alignment=WD_ALIGN_PARAGRAPH.CENTER
set_font(footer.add_run('SGAProjectMain  |  '),8.0)
field=OxmlElement('w:fldSimple'); field.set(qn('w:instr'),'PAGE'); footer._p.append(field)
doc.core_properties.title='SGAProjectMain 구현 구조와 동작 설명서'
doc.core_properties.subject='캐릭터 전투 애니메이션 AI 장비 UI 게임 진행 코드 설명'
doc.core_properties.author='SGAProjectMain 프로젝트 문서'
doc.core_properties.keywords='Unreal Engine 5.4, C++, 구현 설명, 시스템 연결'

p=doc.add_paragraph(style='Title'); inline(p,'SGAProjectMain 구현 구조와 동작 설명서')
title_borders=OxmlElement('w:pBdr')
for edge in ('top','bottom','left','right','between'):
    border=OxmlElement('w:'+edge); border.set(qn('w:val'),'nil'); title_borders.append(border)
p._p.get_or_add_pPr().append(title_borders)
p=doc.add_paragraph(style='Subtitle'); inline(p,'캐릭터 전투 애니메이션 AI 장비 UI 게임 진행')
p=doc.add_paragraph(); inline(p,'2026년 10월 10일  현재 구현 기준')
p=doc.add_paragraph(); inline(p,'이 설명서는 오래된 구현을 다시 읽고 각 시스템의 역할과 호출 연결을 되짚기 위한 자료다. 입력과 선택이 캐릭터 행동, 애니메이션, 장비 사용, 적 AI, 임무 결과와 저장으로 이어지는 과정을 설명한다.')
p=doc.add_paragraph(); inline(p,'본인 담당인 캐릭터·애니메이션·상태·부위 파괴·AI·랜드스케이프와 팀원 담당인 장비·UI·스트라타젬을 함께 다룬다. 주요 클래스와 필드, 데이터 전달, 객체 수명, 대표 소스 위치와 변경 영향이 각 장에 연결되어 있다.')
p=doc.add_paragraph(); p.add_run().add_picture(str(fig_overview),width=Cm(17))
p=doc.add_paragraph(); inline(p,'그림 1  선택에서 전투와 영속 기록까지 이어지는 전체 연결',size=9)
p=doc.add_paragraph(); inline(p,'읽는 순서는 전체 구조 → 캐릭터와 전투 → AI와 전장 → 장비와 지원 호출 → 게임 진행과 UI → 실행 시나리오다. 다시 읽을 때는 목차에서 원하는 기능을 선택하거나 마지막 장의 변경 영향표를 시작점으로 삼는다.')

headings=[]
blocks=[]
for chapter,(fname,label) in enumerate(zip(SOURCES,CHAPTERS),1):
    lines=(NOTES/fname).read_text(encoding='utf-8-sig').splitlines()
    subsection=0; subsub=0; first=True; i=0; code=False; code_lines=[]
    while i<len(lines):
        line=lines[i].rstrip(); i+=1
        if line.startswith('```'):
            if code:
                blocks.append(('code','\n'.join(code_lines),chapter)); code_lines=[]
            code=not code; continue
        if code: code_lines.append(line); continue
        if not line.strip(): continue
        if line.startswith('# '):
            if not first: continue
            first=False
            title=f'{chapter} {label}'; anchor=f'ch_{chapter}'
            headings.append((1,title,anchor)); blocks.append(('h1',(title,anchor),chapter)); continue
        if line.startswith('## '):
            subsection+=1; subsub=0
            title=f'{chapter} {subsection} {clean_heading(line[3:])}'; anchor=f'ch_{chapter}_s_{subsection}'
            headings.append((2,title,anchor)); blocks.append(('h2',(title,anchor),chapter)); continue
        if line.startswith('### '):
            subsub+=1
            title=f'{chapter} {subsection} {subsub} {clean_heading(line[4:])}'
            blocks.append(('h3',title,chapter)); continue
        if line.startswith('|'):
            rows=[line]
            while i<len(lines) and lines[i].startswith('|'):
                rows.append(lines[i]); i+=1
            parsed=[]
            for row in rows:
                vals=[v.strip() for v in row.strip().strip('|').split('|')]
                if all(re.fullmatch(r':?-+:?',v.replace(' ','')) for v in vals): continue
                parsed.append(vals)
            blocks.append(('table',parsed,chapter)); continue
        if re.match(r'^\d+[.)]\s+',line): blocks.append(('number',re.sub(r'^\d+[.)]\s+','',line),chapter)); continue
        if line.startswith('- ') or line.startswith('* '): blocks.append(('bullet',line[2:],chapter)); continue
        if line.startswith('> '): line=line[2:]
        blocks.append(('p',line,chapter))

doc.add_page_break()
p=doc.add_paragraph(); inline(p,'목차',size=17); p.paragraph_format.space_after=Pt(14)
toc_paras=[]
cached_toc=json.loads((WORK/'toc_pages.json').read_text(encoding='utf-8')) if (WORK/'toc_pages.json').exists() else {}
for level,title,anchor in headings:
    p=doc.add_paragraph()
    pf=p.paragraph_format; pf.line_spacing=1.04; pf.space_after=Pt(3.5 if level==2 else 6)
    pf.space_before=Pt(7 if level==1 else 0); pf.keep_with_next=level==1
    if level==2: pf.left_indent=Cm(.4)
    toc_link(p,title,anchor,9.1 if level==2 else 10.5,level==1)
    pf.tab_stops.add_tab_stop(Cm(17),WD_TAB_ALIGNMENT.RIGHT,WD_TAB_LEADER.DOTS)
    set_font(p.add_run('\t'),9.1 if level==2 else 10.5)
    field=OxmlElement('w:fldSimple'); field.set(qn('w:instr'),f'PAGEREF {anchor} \\h')
    field.set(qn('w:dirty'),'true')
    run=OxmlElement('w:r'); text=OxmlElement('w:t'); text.text=str(cached_toc.get(anchor,'')); run.append(text); field.append(run); p._p.append(field)
    toc_paras.append((p,title,anchor,level))

figures={2:(fig_character,'그림 2  입력 상태 애니메이션과 동작 알림의 왕복 연결'),3:(fig_ai,'그림 3  개별 AI 판단과 분대 실행 수명'),4:(fig_gear,'그림 4  장비 생성 사용과 피해 전달')}
for bi,(kind,value,chapter) in enumerate(blocks):
    if kind in ('h1','h2'):
        title,anchor=value
        p=doc.add_paragraph(style='Heading 1' if kind=='h1' else 'Heading 2'); inline(p,title)
        bookmark(p,anchor,bi+100)
        if kind=='h1' and chapter in figures:
            path,caption=figures[chapter]
            p=doc.add_paragraph(); p.add_run().add_picture(str(path),width=Cm(17))
            p.paragraph_format.keep_with_next=True
            p=doc.add_paragraph(); inline(p,caption,size=9); p.paragraph_format.space_after=Pt(10)
    elif kind=='h3':
        p=doc.add_paragraph(style='Heading 3'); inline(p,value)
    elif kind=='table': make_table(doc,value)
    elif kind=='code':
        for line in value.splitlines():
            p=doc.add_paragraph(); p.paragraph_format.space_after=Pt(0); p.paragraph_format.line_spacing=1.08
            r=p.add_run(line); set_font(r,8.7,name='Consolas'); r._element.rPr.rFonts.set(qn('w:eastAsia'),FONT)
    else:
        p=doc.add_paragraph(style='List Bullet' if kind=='bullet' else 'List Number' if kind=='number' else None)
        inline(p,value)

update_fields=OxmlElement('w:updateFields'); update_fields.set(qn('w:val'),'true'); doc.settings.element.append(update_fields)
doc.save(OUT)
manifest={'output':str(OUT),'chapters':CHAPTERS,'headings':[{'level':l,'title':t,'anchor':a} for l,t,a in headings], 'source_chars':sum(len((NOTES/f).read_text(encoding='utf-8-sig')) for f in SOURCES),'paragraphs':len(doc.paragraphs),'tables':len(doc.tables),'figures':4}
(WORK/'manual_manifest.json').write_text(json.dumps(manifest,ensure_ascii=False,indent=2),encoding='utf-8')
print(json.dumps({k:v for k,v in manifest.items() if k!='headings'},ensure_ascii=False))
